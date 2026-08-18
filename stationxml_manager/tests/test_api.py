from __future__ import annotations

import importlib

from app.catalog import seed_catalog
from app.crud import import_hierarchy
from tests.inventories import import_inventory, inventory_with_pz
from tests.test_core import _sample_hierarchy


def test_invalid_xml_returns_korean_400(api_client):
    client, _ = api_client
    response = client.post(
        "/api/import",
        files={"file": ("broken.xml", b"<not-xml", "application/xml")},
        data={"replace_all": "false"},
    )
    assert response.status_code == 400
    assert "파일을 읽지 못했습니다" in response.json()["detail"]


def test_replace_all_requires_explicit_confirmation(api_client):
    client, _ = api_client
    response = client.post(
        "/api/import",
        files={"file": ("inventory.xml", b"<x/>", "application/xml")},
        data={"replace_all": "true", "confirm_replace": "false"},
    )
    assert response.status_code == 400
    assert "전체 교체 확인" in response.json()["detail"]


def test_duplicate_station_returns_conflict(api_client):
    client, SessionLocal = api_client
    session = SessionLocal()
    try:
        seed_catalog(session)
        import_hierarchy(
            session,
            _sample_hierarchy(),
            replace_all=False,
            source="ui",
            actor=None,
        )
        station = session.query(importlib.import_module("app.models").Station).one()
        network_id = station.network_id
    finally:
        session.close()

    response = client.post(
        "/api/stations",
        json={
            "network_id": network_id,
            "code": "AAA",
            "latitude": 37.5,
            "longitude": 127.1,
            "elevation": 80,
        },
    )
    assert response.status_code == 409


def test_template_endpoint_contains_no_inventory(api_client):
    client, SessionLocal = api_client
    session = SessionLocal()
    try:
        seed_catalog(session)
        import_hierarchy(
            session,
            _sample_hierarchy(),
            replace_all=False,
            source="ui",
            actor=None,
        )
    finally:
        session.close()

    response = client.get("/api/template.xlsx")
    assert response.status_code == 200
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(response.content))
    assert workbook["channels"].max_row == 1


def test_api_key_error_keeps_cors_headers(api_client, monkeypatch):
    client, _ = api_client
    main_module = importlib.import_module("app.main")
    monkeypatch.setattr(main_module, "_api_key", "secret")
    response = client.get(
        "/api/health",
        headers={
            "Origin": "http://localhost:5173",
            "X-API-Key": "wrong",
        },
    )
    assert response.status_code == 401
    assert response.headers["access-control-allow-origin"] == "http://localhost:5173"

    ok = client.get("/api/health", headers={"X-API-Key": "secret"})
    assert ok.status_code == 200


def test_api_create_custom_catalog_without_code(api_client):
    client, SessionLocal = api_client
    session = SessionLocal()
    try:
        seed_catalog(session)
    finally:
        session.close()
    response = client.post(
        "/api/catalog",
        json={"kind": "sensor", "manufacturer": "Acme", "model": "Geophone"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["code"] == "CUSTOM_Acme_Geophone"
    assert body["origin"] == "custom"


def test_export_status_empty_and_ready(api_client):
    client, SessionLocal = api_client
    empty = client.get("/api/export/status")
    assert empty.status_code == 200
    body = empty.json()
    assert body["channel_count"] == 0
    assert body["stationxml_ok"] is False
    assert body["dataless_ok"] is False

    session = SessionLocal()
    try:
        import_inventory(session, inventory_with_pz(site_name="AAA"))
    finally:
        session.close()
    ready = client.get("/api/export/status")
    assert ready.status_code == 200
    ok = ready.json()
    assert ok["channel_count"] == 1
    assert ok["response_count"] == 1
    assert ok["stationxml_ok"] is True
    assert ok["dataless_ok"] is True


def test_help_page_covers_metadata_methods():
    from pathlib import Path

    text = Path("frontend/src/HelpPage.tsx").read_text(encoding="utf-8")
    assert "메타데이터 생성 방법" in text
    assert "엑셀로 만들기" in text
    assert "StationXML" in text
    assert "Dataless SEED" in text
    assert "NRL" in text
    assert "Poles/Zeros" in text
