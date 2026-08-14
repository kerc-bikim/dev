from __future__ import annotations

from io import BytesIO

import pytest
from obspy import UTCDateTime
from obspy.core.inventory import Channel, Inventory, Network, Station
from obspy.core.inventory.response import InstrumentSensitivity, Response
from obspy.core.inventory.util import Equipment, Site
from sqlalchemy.orm import Session, sessionmaker

from app.catalog import find_by_manufacturer_model, seed_catalog
from app.columns import resolve_header
from app.crud import (
    create_catalog_item,
    delete_catalog_item,
    export_stationxml_bytes,
    import_hierarchy,
    list_channels,
    update_catalog_item,
    update_channel,
    update_station,
)
from app.db import Base, make_engine
from app.errors import ValidationError
from app.excel_io import read_excel, write_excel, write_template
from app.validation import (
    infer_az_dip,
    validate_lat_lon,
    validate_sample_rate,
    validate_time_order,
)
from app.xml_io import read_stationxml


@pytest.fixture
def session(tmp_path) -> Session:
    engine = make_engine(f"sqlite:///{tmp_path}/test.db")
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False)
    s = SessionLocal()
    seed_catalog(s)
    yield s
    s.close()


def test_header_aliases():
    assert resolve_header("위도") == "latitude"
    assert resolve_header("센서ID") == "sensor_id"
    assert resolve_header("unknown") is None


def test_validation_helpers():
    validate_lat_lon(37.5, 127.0)
    with pytest.raises(ValidationError):
        validate_lat_lon(120, 127)
    validate_sample_rate(100)
    with pytest.raises(ValidationError):
        validate_sample_rate(0)
    validate_time_order(UTCDateTime(2020, 1, 1), UTCDateTime(2021, 1, 1), "시작", "끝")
    with pytest.raises(ValidationError):
        validate_time_order(
            UTCDateTime(2021, 1, 1), UTCDateTime(2020, 1, 1), "시작", "끝"
        )
    assert infer_az_dip("HHZ") == (0.0, -90.0)
    assert infer_az_dip("HHE") == (90.0, 0.0)


def _sample_hierarchy():
    return {
        "warnings": [],
        "catalog": {"sensors": [], "dataloggers": []},
        "networks": {
            "XX": {
                "code": "XX",
                "description": "테스트망",
                "operator_agency": "테스트",
                "restricted_status": "open",
                "stations": {
                    "AAA": {
                        "code": "AAA",
                        "latitude": 37.5,
                        "longitude": 127.1,
                        "elevation": 80.0,
                        "site_name": "첫번째",
                        "site_description": None,
                        "site_town": None,
                        "site_region": None,
                        "site_country": "KR",
                        "vault": None,
                        "geology": None,
                        "description": None,
                        "creation_date": "2020-01-01T00:00:00.000000Z",
                        "termination_date": None,
                        "channels": [
                            {
                                "location": "",
                                "channel": "HHZ",
                                "start_time": "2020-01-01T00:00:00.000000Z",
                                "end_time": None,
                                "sample_rate": 100.0,
                                "depth": 0.0,
                                "azimuth": 0.0,
                                "dip": -90.0,
                                "description": None,
                                "comment": "비고",
                                "channel_types": "GEOPHYSICAL,CONTINUOUS",
                                "clock_drift": None,
                                "latitude": None,
                                "longitude": None,
                                "elevation": None,
                                "sensor_id": "Guralp_CMG-3T",
                                "sensor_serial": "S1",
                                "sensor_type": None,
                                "sensor_install_date": None,
                                "sensor_remove_date": None,
                                "datalogger_id": "REFTEK_RT130_100sps",
                                "datalogger_serial": "D1",
                                "datalogger_type": None,
                                "datalogger_install_date": None,
                                "datalogger_remove_date": None,
                            }
                        ],
                    }
                },
            }
        },
    }


def test_import_and_station_edit(session):
    result = import_hierarchy(
        session, _sample_hierarchy(), replace_all=False, source="ui", actor="테스터"
    )
    assert result["created"] == 1
    ch = list_channels(session)[0]
    assert ch.station.site_name == "첫번째"
    update_station(session, ch.station.id, {"site_name": "수정됨"}, "테스터")
    ch = list_channels(session)[0]
    assert ch.station.site_name == "수정됨"


def test_bad_sensor_id(session):
    data = _sample_hierarchy()
    data["networks"]["XX"]["stations"]["AAA"]["channels"][0]["sensor_id"] = "NOPE"
    with pytest.raises(ValidationError, match="알 수 없는 센서ID"):
        import_hierarchy(session, data, replace_all=False, source="excel", actor=None)


def test_sample_rate_mismatch(session):
    data = _sample_hierarchy()
    data["networks"]["XX"]["stations"]["AAA"]["channels"][0]["sample_rate"] = 200
    with pytest.raises(ValidationError, match="샘플링레이트"):
        import_hierarchy(session, data, replace_all=False, source="excel", actor=None)


def test_catalog_delete_in_use(session):
    import_hierarchy(
        session, _sample_hierarchy(), replace_all=False, source="ui", actor=None
    )
    from app.models import EquipmentCatalog

    row = session.query(EquipmentCatalog).filter_by(code="Guralp_CMG-3T").one()
    with pytest.raises(ValidationError, match="사용하는 채널"):
        delete_catalog_item(session, row.id, None)


def test_catalog_create_and_delete(session):
    item = create_catalog_item(
        session,
        {
            "kind": "sensor",
            "code": "Test_Sensor",
            "manufacturer": "TestCo",
            "model": "T1",
        },
        "테스터",
    )
    delete_catalog_item(session, item.id, "테스터")


def test_conflicting_site_name_excel(session, tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "bad.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "channels"
    ws.append(
        [
            "네트워크",
            "관측소",
            "채널",
            "위도",
            "경도",
            "시작시간",
            "샘플링레이트",
            "관측소명",
        ]
    )
    ws.append(["XX", "AAA", "HHZ", 37.5, 127.0, "2020-01-01", 100, "A"])
    ws.append(["XX", "AAA", "HHN", 37.5, 127.0, "2020-01-01", 100, "B"])
    wb.save(path)
    with pytest.raises(ValidationError, match="관측소명"):
        read_excel(path)


def test_unknown_excel_header(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "bad.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "네트워크",
            "관측소",
            "채널",
            "위도",
            "경도",
            "시작시간",
            "샘플링레이트",
            "담당자",
        ]
    )
    ws.append(["XX", "AAA", "HHZ", 37.5, 127.0, "2020-01-01", 100, "홍길동"])
    wb.save(path)
    with pytest.raises(ValidationError, match="알 수 없는 엑셀 열"):
        read_excel(path)


def test_excel_reimport_keeps_response(session):
    inv = _inventory_with_response()
    buf = BytesIO()
    inv.write(buf, format="STATIONXML")
    buf.seek(0)
    hierarchy = read_stationxml(buf, session)
    import_hierarchy(session, hierarchy, replace_all=False, source="xml", actor=None)
    ch = list_channels(session)[0]
    assert ch.response_xml
    original = ch.response_xml

    xlsx = write_excel(session)
    excel_h = read_excel(BytesIO(xlsx))
    import_hierarchy(session, excel_h, replace_all=False, source="excel", actor=None)
    ch = list_channels(session)[0]
    assert ch.response_xml == original


def test_xml_response_roundtrip(session):
    inv = _inventory_with_response()
    buf = BytesIO()
    inv.write(buf, format="STATIONXML")
    buf.seek(0)
    hierarchy = read_stationxml(buf, session)
    import_hierarchy(session, hierarchy, replace_all=True, source="xml", actor=None)
    update_station(
        session, list_channels(session)[0].station.id, {"site_name": "변경"}, None
    )
    exported = export_stationxml_bytes(session)
    from obspy import read_inventory

    back = read_inventory(BytesIO(exported), format="STATIONXML")
    cha = back[0][0][0]
    assert cha.response is not None
    assert cha.sensor.manufacturer == "Guralp"


def test_depth_zero_channel_update(session):
    import_hierarchy(
        session, _sample_hierarchy(), replace_all=False, source="ui", actor=None
    )
    ch = list_channels(session)[0]
    updated = update_channel(
        session,
        ch.id,
        {
            "depth": 0,
            "comment": "수정",
            "sample_rate": 100,
            "sensor_id": ch.sensor_id,
            "datalogger_id": ch.datalogger_id,
        },
        None,
    )
    assert updated.depth == 0
    assert updated.comment == "수정"


def test_excel_time_normalization_does_not_duplicate(session):
    import_hierarchy(
        session, _sample_hierarchy(), replace_all=False, source="ui", actor=None
    )
    excel_hierarchy = read_excel(BytesIO(write_excel(session)))
    import_hierarchy(
        session,
        excel_hierarchy,
        replace_all=False,
        source="excel",
        actor=None,
    )
    assert len(list_channels(session)) == 1


def test_replace_all_excel_keeps_matching_response(session):
    inv = _inventory_with_response()
    buf = BytesIO()
    inv.write(buf, format="STATIONXML")
    buf.seek(0)
    import_hierarchy(
        session,
        read_stationxml(buf, session),
        replace_all=False,
        source="xml",
        actor=None,
    )
    response_before = list_channels(session)[0].response_xml
    excel_hierarchy = read_excel(BytesIO(write_excel(session)))
    import_hierarchy(
        session,
        excel_hierarchy,
        replace_all=True,
        source="excel",
        actor=None,
    )
    rows = list_channels(session)
    assert len(rows) == 1
    assert rows[0].response_xml == response_before


def test_replace_all_rejects_empty_inventory(session):
    import_hierarchy(
        session, _sample_hierarchy(), replace_all=False, source="ui", actor=None
    )
    with pytest.raises(ValidationError, match="가져올 채널이 없습니다"):
        import_hierarchy(
            session,
            {"networks": {}, "catalog": {}, "warnings": []},
            replace_all=True,
            source="xml",
            actor=None,
        )
    session.rollback()
    assert len(list_channels(session)) == 1


def test_template_does_not_include_inventory_rows(session):
    import_hierarchy(
        session, _sample_hierarchy(), replace_all=False, source="ui", actor=None
    )
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(write_template(session)))
    assert workbook["channels"].max_row == 1
    assert workbook["catalog_sensors"].max_row > 1


def test_station_can_move_to_another_network(session):
    import_hierarchy(
        session, _sample_hierarchy(), replace_all=False, source="ui", actor=None
    )
    from app.models import Network

    target = Network(code="YY")
    session.add(target)
    session.commit()
    station = list_channels(session)[0].station
    update_station(session, station.id, {"network_id": target.id}, None)
    assert list_channels(session)[0].station.network.code == "YY"


def test_datalogger_rate_mismatch_does_not_fallback(session):
    match = find_by_manufacturer_model(
        session,
        "datalogger",
        "REF TEK",
        "RT 130",
        999,
    )
    assert match is None


def test_catalog_import_does_not_erase_existing_validation_fields(session):
    from app.models import EquipmentCatalog

    sensor = session.query(EquipmentCatalog).filter_by(code="Guralp_CMG-3T").one()
    logger = session.query(EquipmentCatalog).filter_by(code="REFTEK_RT130_100sps").one()
    original_sensor_keys = sensor.nrl_keys
    original_rate = logger.sample_rate
    hierarchy = _sample_hierarchy()
    hierarchy["catalog"] = {
        "sensors": [
            {
                "code": sensor.code,
                "manufacturer": sensor.manufacturer,
                "model": sensor.model,
                "nrl_keys": None,
            }
        ],
        "dataloggers": [
            {
                "code": logger.code,
                "manufacturer": logger.manufacturer,
                "model": logger.model,
                "sample_rate": None,
                "nrl_keys": None,
            }
        ],
    }
    import_hierarchy(session, hierarchy, replace_all=False, source="excel", actor=None)
    session.refresh(sensor)
    session.refresh(logger)
    assert sensor.nrl_keys == original_sensor_keys
    assert logger.sample_rate == original_rate


def test_legacy_start_time_matches_canonical_import(session):
    hierarchy = _sample_hierarchy()
    hierarchy["networks"]["XX"]["stations"]["AAA"]["channels"][0]["start_time"] = (
        "2020-01-01T00:00:00"
    )
    import_hierarchy(session, hierarchy, replace_all=False, source="ui", actor=None)
    assert len(list_channels(session)) == 1

    canonical = _sample_hierarchy()
    import_hierarchy(session, canonical, replace_all=False, source="excel", actor=None)
    rows = list_channels(session)
    assert len(rows) == 1
    assert rows[0].start_time == "2020-01-01T00:00:00.000000Z"


def test_catalog_import_rejects_non_positive_sample_rate(session):
    hierarchy = _sample_hierarchy()
    hierarchy["catalog"] = {
        "sensors": [],
        "dataloggers": [
            {
                "code": "NEW_BAD_LOGGER",
                "manufacturer": "Test",
                "model": "Bad",
                "sample_rate": 0,
                "nrl_keys": None,
            }
        ],
    }
    with pytest.raises(ValidationError, match="0보다 커야"):
        import_hierarchy(
            session, hierarchy, replace_all=False, source="excel", actor=None
        )


def test_used_datalogger_sample_rate_cannot_be_cleared(session):
    import_hierarchy(
        session, _sample_hierarchy(), replace_all=False, source="ui", actor=None
    )
    from app.models import EquipmentCatalog

    logger = session.query(EquipmentCatalog).filter_by(code="REFTEK_RT130_100sps").one()
    with pytest.raises(ValidationError, match="비울 수 없습니다"):
        update_catalog_item(session, logger.id, {"sample_rate": None}, None)


def _inventory_with_response() -> Inventory:
    resp = Response(
        instrument_sensitivity=InstrumentSensitivity(
            value=1.5e9,
            frequency=1.0,
            input_units="M/S",
            output_units="COUNTS",
        )
    )
    cha = Channel(
        code="HHZ",
        location_code="",
        latitude=37.5,
        longitude=127.1,
        elevation=80.0,
        depth=0.0,
        azimuth=0.0,
        dip=-90.0,
        sample_rate=100.0,
        start_date=UTCDateTime(2020, 1, 1),
        sensor=Equipment(manufacturer="Guralp", model="CMG-3T", serial_number="S1"),
        data_logger=Equipment(
            manufacturer="REF TEK", model="RT 130", serial_number="D1"
        ),
    )
    cha.response = resp
    sta = Station(
        code="AAA",
        latitude=37.5,
        longitude=127.1,
        elevation=80.0,
        creation_date=UTCDateTime(2020, 1, 1),
        site=Site(name="첫번째"),
        channels=[cha],
    )
    net = Network(code="XX", stations=[sta], description="테스트망")
    return Inventory(networks=[net], source="test")
